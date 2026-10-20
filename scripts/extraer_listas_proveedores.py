"""Genera data/proveedores_listas.json a partir del libro «Entity - Falcon Coffees - Perú - Legacy.xlsx».

Solo extrae listas de referencia (países, tipos de persona y de documento, ubigeos y códigos postales);
no copia datos de proveedores. Uso:  python scripts/extraer_listas_proveedores.py ruta/al/libro.xlsx
"""
import json
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

SALIDA = Path(__file__).resolve().parent.parent / "data" / "proveedores_listas.json"


def _col(ws, col, desde, hasta):
    return [str(v).strip() for v in (ws.cell(r, col).value for r in range(desde, hasta + 1)) if v not in (None, "")]


def main(ruta: str) -> None:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(ruta, data_only=True, read_only=False)
    valid, no_tocar_valid = wb["DATA VALID"], wb["DATA VALID NO TOCAR"]
    cp = wb["Codigos Postales"]
    zips: dict[str, str] = {}
    for fila in cp.iter_rows(min_row=3, values_only=True):  # clave (col I) -> código postal (col H), primera coincidencia
        clave, codigo = fila[8], fila[7]
        if clave and codigo and clave not in zips:
            zips[str(clave).strip()] = str(codigo).strip()
    datos = {
        "paises": _col(valid, 11, 3, 258),                       # Vendor!I = 'DATA VALID'!K3:K258
        "personas": _col(no_tocar_valid, 11, 275, 277),           # Vendor!O
        "tipos_doc": _col(no_tocar_valid, 10, 282, 289),          # Vendor!P
        "ubigeos": [str(c[0].value).strip() for c in wb["No tocar"].iter_rows(min_col=1, max_col=1) if c[0].value],  # Vendor!K
        "zip": zips,
    }
    SALIDA.write_text(json.dumps(datos, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print({k: len(v) for k, v in datos.items()}, "->", SALIDA)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
