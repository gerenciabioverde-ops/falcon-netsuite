"""Carga de un Excel/CSV a las grillas: busca la hoja, detecta la fila de encabezados y mapea por nombre.

Sirve tanto para el Excel de plantillas (una hoja por pestaña, encabezados en la fila 1) como para el libro de
proveedores (hoja «Vendor», encabezados en la fila 3 y filas de ayuda/ejemplo antes de los datos).
"""
from __future__ import annotations

import csv
import io
import warnings
from html import unescape
from html.parser import HTMLParser
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from .esquema import Hoja, Plantilla, norm

MAX_FILAS = 2000
ALIAS_HOJAS = {"vendor": "proveedores", "vendors": "proveedores", "proveedor": "proveedores"}
# Etiquetas de la primera columna de las filas de ayuda del libro de carga (no son datos)
FILAS_DE_AYUDA = {"totales", "total", "campos", "oracle netsuite standard", "tipo de campo", "mandatory", "mandatoriedad", "description",
                  "descripcion", "descripcción", "ejemplo", "comentarios", "fields"}


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, bool):
        return "T" if v else "F"
    if isinstance(v, float) or isinstance(v, int):
        try:
            return format(Decimal(str(v)).normalize(), "f")
        except InvalidOperation:
            return str(v)
    return str(v).strip()


# nombres con que suelen venir ciertas columnas del reporte -> nombre de la columna de la grilla
ALIAS_COLUMNAS = {"nro operacion": "n° operación", "nro. operacion": "n° operación", "n° operacion": "n° operación",
                  "nro operación": "n° operación", "numero de operacion": "n° operación", "número de operación": "n° operación"}


def _encabezados(filas: list[list], hoja: Hoja) -> tuple[int, dict[int, int]] | None:
    """(índice de la fila de encabezados, {columna de la grilla: columna del archivo}) o None."""
    objetivos = [norm(e) for e in hoja.encabezados]
    mejor: tuple[int, dict[int, int]] | None = None
    for i, fila in enumerate(filas[:25]):
        presentes = [ALIAS_COLUMNAS.get(norm(c), norm(c)) for c in fila]
        usados: set[int] = set()
        mapa: dict[int, int] = {}
        for g, nombre in enumerate(objetivos):  # duplicados (p. ej. «Memo» dos veces): se asignan en orden
            for j, cabecera in enumerate(presentes):
                if cabecera == nombre and j not in usados:
                    mapa[g] = j
                    usados.add(j)
                    break
        if len(mapa) >= min(3, max(1, len(objetivos) - 1)) and (mejor is None or len(mapa) > len(mejor[1])):
            mejor = (i, mapa)
    return mejor


def _numero(t: str):
    try:
        return Decimal(str(t).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def _completar_bancos(cabecera: list[list], filas: list[list[str]], hoja: Hoja) -> None:
    """Reporte del banco: entidad y cuenta salen del encabezado del archivo; la columna Saldos viene como fórmula (=(Q8+O9)) y se
    calcula aquí: saldo anterior + ingreso − egreso (la fila del saldo inicial trae su valor)."""
    datos = {}
    for f in cabecera:
        if len(f) >= 2 and norm(f[0]).rstrip(":") in ("entidad financiera", "código de la cuenta", "codigo de la cuenta"):
            datos["entidad" if norm(f[0]).startswith("entidad") else "cuenta"] = f[1].strip()
    i_ent, i_cta = hoja.idx("Entidad Financiera"), hoja.idx("Cuenta bancaria")
    i_ing, i_egr, i_sal = hoja.idx("Ingreso"), hoja.idx("Egreso"), hoja.idx("Saldos")
    saldo = Decimal(0)
    for c in filas:
        c[i_ent], c[i_cta] = datos.get("entidad", ""), datos.get("cuenta", "")
        actual = _numero(c[i_sal]) if not c[i_sal].startswith("=") else None
        if actual is None:
            saldo = saldo + (_numero(c[i_ing]) or 0) - (_numero(c[i_egr]) or 0)
            c[i_sal] = format(saldo.quantize(Decimal("0.01")), "f")
        else:
            saldo = actual


def _completar_liquidaciones(cabecera: list, filas: list[list]) -> list[list]:
    """Reporte de liquidaciones: cuando una L.C. agrupa varias recepciones, el reporte omite las celdas combinadas (Serie y N° COMP. y
    S/ TOTAL L.C.) en las filas siguientes y las columnas se corren. Aquí se vuelven a insertar y la serie se copia de la fila anterior."""
    nombres = [norm(c) for c in cabecera]
    i_ser = nombres.index("serie y n° comp.") if "serie y n° comp." in nombres else -1
    i_tot = next((i for i, n in enumerate(nombres) if n.startswith("s/ total l.c")), -1)
    if i_ser < 0 or i_tot < 0:
        return filas
    salida, serie = [], ""
    for f in filas:
        f = list(f)
        if len(f) == len(nombres) - 2:
            f.insert(i_ser, serie)
            f.insert(i_tot, "")
        elif len(f) > i_ser:
            serie = str(f[i_ser]).strip()
        salida.append(f)
    return salida


def _filas_de_datos(filas: list[list], hoja: Hoja, avisos: list[str]) -> list[list[str]] | None:
    enc = _encabezados(filas, hoja)
    if enc is None:
        return None
    inicio, mapa = enc
    if hoja.nombre == "Liquidaciones":
        filas = filas[:inicio + 1] + _completar_liquidaciones(filas[inicio], filas[inicio + 1:])
    salida: list[list[str]] = []
    for fila in filas[inicio + 1:]:
        if fila and norm(fila[0]) in FILAS_DE_AYUDA:
            continue
        celdas = [""] * len(hoja.encabezados)
        for g, j in mapa.items():
            if j < len(fila):
                celdas[g] = _texto(fila[j])
        if any(celdas):
            salida.append(celdas)
    if hoja.nombre in ("Bancos MN", "Bancos ME"):
        _completar_bancos(filas[:inicio], salida, hoja)
    if hoja.nombre == "Liquidaciones":  # el encabezado del reporte dice «Almacen: Jaén»; con «Todos» la sede queda vacía (se elige al cargar)
        for f in filas[:inicio]:
            for c in f:
                t = str(c).strip()
                if norm(t).startswith("almacen:") and norm(t.split(":", 1)[1]) not in ("", "todos", "todas"):
                    i_sede = hoja.idx("Sede")
                    for fila in salida:
                        fila[i_sede] = t.split(":", 1)[1].strip()
    if hoja.nombre == "Proveedores":  # «No encontrado» (resultado del XLOOKUP del libro) no es un código postal
        i = hoja.idx("Código ZIP")
        for c in salida:
            if c[i] and not c[i].isdigit():
                c[i] = ""
    if len(salida) > MAX_FILAS:
        avisos.append(f"«{hoja.nombre.strip()}» tiene {len(salida)} filas; se cargaron las primeras {MAX_FILAS}.")
        salida = salida[:MAX_FILAS]
    return salida


class _TablaHtml(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.filas: list[list[str]] = []
        self._fila: list[str] | None = None
        self._celda: str | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._fila = []
        elif tag in ("td", "th"):
            self._celda = ""

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._fila is not None and self._celda is not None:
            self._fila.append(unescape(self._celda).strip())
            self._celda = None
        elif tag == "tr" and self._fila is not None:
            self.filas.append(self._fila)
            self._fila = None

    def handle_data(self, data):
        if self._celda is not None:
            self._celda += data


def _leer_tabla_html(contenido: bytes, plantilla: Plantilla, clave_pestana: str, avisos: list[str]) -> dict[str, list[list[str]]]:
    """Reportes que el sistema exporta como «.xls» pero que en realidad son una tabla HTML."""
    texto = contenido.decode("utf-8-sig", errors="replace")
    if "<table" not in texto.lower():
        raise ValueError("Este .xls es del formato antiguo de Excel. Ábralo en Excel y guárdelo como .xlsx (o como CSV) y vuelva a cargarlo.")
    parser = _TablaHtml()
    parser.feed(texto)
    filas = parser.filas
    candidatas = ([clave_pestana] if clave_pestana in plantilla.hojas else []) + [k for k in plantilla.hojas if k != clave_pestana]
    for clave in candidatas:
        hoja = plantilla.hojas[clave]
        enc = _encabezados(filas, hoja)
        if enc is None or len(enc[1]) < min(len(hoja.encabezados), max(4, len(hoja.encabezados) // 2)):
            continue
        datos = _filas_de_datos(filas, hoja, avisos)
        if datos:
            if clave != clave_pestana:
                avisos.append(f"El reporte se cargó en «{hoja.nombre.strip()}» por sus encabezados.")
            return {clave: datos}
    raise ValueError("No reconocí los encabezados del reporte con los de ninguna pestaña.")


def leer(nombre_archivo: str, contenido: bytes, plantilla: Plantilla, clave_csv: str = "") -> tuple[dict[str, list[list[str]]], list[str]]:
    """Devuelve ({clave de pestaña: filas}, avisos). Lanza ValueError si el archivo no sirve."""
    avisos: list[str] = []
    nombre = nombre_archivo.lower()
    resultado: dict[str, list[list[str]]] = {}

    if nombre.endswith(".csv"):
        if clave_csv not in plantilla.hojas:
            raise ValueError("Para un CSV cárguelo desde la pestaña a la que corresponde.")
        texto = contenido.decode("utf-8-sig", errors="replace")
        primera = texto.split("\n", 1)[0]
        delim = ";" if primera.count(";") > primera.count(",") else ","
        filas = [list(f) for f in csv.reader(io.StringIO(texto), delimiter=delim)]
        datos = _filas_de_datos(filas, plantilla.hojas[clave_csv], avisos)
        if datos is None:
            raise ValueError("No reconocí los encabezados del CSV para esta pestaña.")
        return {clave_csv: datos}, avisos

    if nombre.endswith((".xls", ".html", ".htm")):
        return _leer_tabla_html(contenido, plantilla, clave_csv, avisos), avisos
    if not nombre.endswith((".xlsx", ".xlsm")):
        raise ValueError("Suba un archivo .xlsx, .xls (reporte) o .csv")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    for hoja_xl in wb.sheetnames:
        clave = ALIAS_HOJAS.get(norm(hoja_xl), norm(hoja_xl))
        if clave not in plantilla.hojas or clave in resultado:
            continue
        filas = [list(f) for f in wb[hoja_xl].iter_rows(values_only=True)]
        datos = _filas_de_datos(filas, plantilla.hojas[clave], avisos)
        if datos is None:
            avisos.append(f"La hoja «{hoja_xl.strip()}» no tiene los encabezados esperados; se omitió.")
        elif datos:
            resultado[clave] = datos
    if not resultado:  # ninguna hoja con nombre conocido: se identifican por sus encabezados
        for hoja_xl in wb.sheetnames:
            ws = wb[hoja_xl]
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            muestra = [list(f) for f in ws.iter_rows(min_row=1, max_row=25, values_only=True)]
            for clave, hoja in plantilla.hojas.items():
                enc = _encabezados(muestra, hoja)
                if clave in resultado or enc is None or len(enc[1]) < min(len(hoja.encabezados), max(4, len(hoja.encabezados) // 2)):
                    continue
                datos = _filas_de_datos([list(f) for f in ws.iter_rows(values_only=True)], hoja, avisos)
                if datos:
                    resultado[clave] = datos
                    avisos.append(f"La hoja «{hoja_xl.strip()}» se cargó en «{hoja.nombre.strip()}» por sus encabezados.")
                    break
    if not resultado:
        raise ValueError("No encontré datos con los encabezados de ninguna pestaña. Revise que la fila de encabezados "
                         "tenga los nombres de las columnas (por ejemplo «Nombre de la empresa», «Latam - Tax Reg Number»).")
    return resultado, avisos
