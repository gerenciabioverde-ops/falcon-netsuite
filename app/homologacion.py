"""Cuentas homologadas: equivalencia entre la cuenta contable de Falcon y la cuenta de NetSuite.

La tabla se sube desde la pantalla, se guarda en `datos_usuario/` (fuera de git, para que
«Actualizar» no la pise) y se aplica al generar los CSV en las columnas de cuenta.
"""
from __future__ import annotations

import csv
import io
import json
import re
import time
import warnings
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook

RUTA = Path(__file__).resolve().parent.parent / "datos_usuario" / "cuentas_homologadas.json"
MAX_FILAS = 20000

# Encabezados reconocidos (se comparan sin tildes ni mayúsculas)
_ALIAS_FALCON = ("falcon", "codigo", "pcge", "cuenta contable", "cuenta origen")
_ALIAS_NS = ("netsuite", "id interno", "internal", "cuenta destino")
_ALIAS_NOMBRE = ("nombre", "descripcion", "denominacion")


def _sin_tildes(t: str) -> str:
    return (t.lower().replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u"))


def normalizar(valor) -> str:
    """Forma comparable de una cuenta: sin espacios y sin «.0» de Excel."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"\s+", "", str(valor)).strip()


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _filas_de_archivo(nombre: str, contenido: bytes) -> list[list]:
    nombre = nombre.lower()
    if nombre.endswith(".csv"):
        texto = contenido.decode("utf-8-sig", errors="replace")
        primera = texto.split("\n", 1)[0]
        delim = ";" if primera.count(";") > primera.count(",") else ","
        return [f for f in csv.reader(io.StringIO(texto), delimiter=delim)]
    if nombre.endswith((".xlsx", ".xlsm")):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        return [list(f) for f in ws.iter_rows(values_only=True)]
    raise ValueError("Suba un archivo .xlsx o .csv")


def _columnas(encabezado: list) -> tuple[int, int, int | None, int | None]:
    """Índices de: cuenta Falcon, cuenta NetSuite, nombre Falcon, nombre NetSuite."""
    h = [_sin_tildes(_texto(c)) for c in encabezado]
    i_ns = next((i for i, t in enumerate(h) if any(a in t for a in _ALIAS_NS) and not t.startswith("nombre")), None)
    i_fa = next((i for i, t in enumerate(h) if any(a in t for a in _ALIAS_FALCON)
                 and "netsuite" not in t and not t.startswith("nombre")), None)
    if i_fa is None or i_ns is None:
        raise ValueError("No encontré las columnas. La primera fila debe tener «Cuenta Falcon» y «ID interno NetSuite» "
                         "(opcionalmente «Nombre Falcon» y «Nombre NetSuite»). Descargue la plantilla.")
    n_fa = next((i for i, t in enumerate(h) if any(a in t for a in _ALIAS_NOMBRE) and "falcon" in t), None)
    n_ns = next((i for i, t in enumerate(h) if any(a in t for a in _ALIAS_NOMBRE) and "netsuite" in t), None)
    return i_fa, i_ns, n_fa, n_ns


def _num_cuenta(v) -> str:
    """Número de cuenta NetSuite comparable: 1100 == '1100.00', 1045.1 == '1045.10'."""
    t = _texto(v)
    try:
        return format(Decimal(t).normalize(), "f")
    except InvalidOperation:
        return t


def _encabezado(filas: list[list], requeridos: tuple[str, ...], limite: int = 20):
    for i, f in enumerate(filas[:limite]):
        h = [_sin_tildes(_texto(c)) for c in f]
        if all(any(r in t for t in h) for r in requeridos):
            return i, h
    return None


def _desde_coa_pcge(nombre: str, contenido: bytes) -> tuple[list[dict], list[str], list[str]] | None:
    """Formato del COA de NetSuite + hoja de equivalencia con el PCGE (Peru Tax Code).

    Cruza: código PCGE -> número de cuenta NetSuite -> ID interno (hoja COA). None si el archivo no es de este formato.
    """
    if not nombre.lower().endswith((".xlsx", ".xlsm")):
        return None
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    hojas = {n: [list(f) for f in wb[n].iter_rows(values_only=True)] for n in wb.sheetnames}
    coa = pcge = None
    for filas in hojas.values():
        if coa is None and (e := _encabezado(filas, ("id interno", "numero"))):
            coa = (filas, *e)
        if pcge is None and (e := _encabezado(filas, ("coa - netsuite", "peru tax code"))):
            pcge = (filas, *e)
    if not (coa and pcge):
        return None

    col = lambda h, igual=None, contiene=None: next(  # noqa: E731
        (i for i, t in enumerate(h) if (igual and t == igual) or (contiene and contiene in t)), None)
    f_coa, e_coa, h_coa = coa
    i_id, i_num, i_cta = col(h_coa, contiene="id interno"), col(h_coa, igual="numero"), col(h_coa, igual="cuenta")
    i_inact, i_res = col(h_coa, igual="inactivo"), col(h_coa, igual="resumen")
    if i_id is None or i_num is None:
        return None
    plan: dict[str, list] = {}
    for f in f_coa[e_coa + 1:]:
        f = list(f) + [None] * 12
        if _texto(f[i_num]):
            plan[_num_cuenta(f[i_num])] = f

    f_pe, e_pe, h_pe = pcge
    i_ns, i_pe = col(h_pe, contiene="coa - netsuite"), col(h_pe, contiene="peru tax code")
    i_desc = col(h_pe, igual="summary description")
    i_dest = col(h_pe, contiene="cuentas destino")  # la contrapartida va en la columna siguiente (sin encabezado)
    i_contra = i_dest + 1 if i_dest is not None else None
    cuentas: dict[str, dict] = {}
    errores: list[str] = []
    avisos: list[str] = []
    for nro, f in enumerate(f_pe[e_pe + 1:], start=e_pe + 2):
        f = list(f) + [None] * 12
        ns, pe = _texto(f[i_ns]), _texto(f[i_pe])
        if not ns and not pe:
            continue
        if not pe:
            avisos.append(f"Fila {nro}: la cuenta NetSuite {ns} no tiene código PCGE; se omitió.")
            continue
        if not ns:
            avisos.append(f"Fila {nro}: el PCGE {pe} no tiene cuenta NetSuite; se omitió.")
            continue
        fila_plan = plan.get(_num_cuenta(ns))
        if fila_plan is None:
            avisos.append(f"Fila {nro}: la cuenta NetSuite {ns} (PCGE {pe}) no existe en el COA; se omitió.")
            continue
        id_interno = _texto(fila_plan[i_id])
        previa = cuentas.get(normalizar(pe))
        if previa and previa["netsuite"] != id_interno:
            errores.append(f"Fila {nro}: el PCGE {pe} ya estaba asignado a {previa['netsuite']} y ahora a {id_interno}.")
            continue
        if (i_inact is not None and _texto(fila_plan[i_inact]).lower() in ("sí", "si")) or \
           (i_res is not None and _texto(fila_plan[i_res]).lower() in ("sí", "si")):
            avisos.append(f"PCGE {pe} → cuenta NetSuite {ns} está inactiva o es de resumen: NetSuite no permitirá registrar en ella.")
        cuentas[normalizar(pe)] = {
            "falcon": pe,
            "netsuite": id_interno,
            "nombre_falcon": _texto(f[i_desc]) if i_desc is not None else "",
            "nombre_netsuite": (ns + " " + _texto(fila_plan[i_cta])).strip() if i_cta is not None else ns,
            # asiento de destino del PCGE: solo informativo (se hace en Falcon Contable, no en NetSuite)
            "destino": _texto(f[i_dest]) if i_dest is not None else "",
            "contrapartida": _texto(f[i_contra]) if i_contra is not None else "",
        }
    if not cuentas:
        raise ValueError("No se pudo cruzar ninguna cuenta entre el COA y la hoja de equivalencias.")
    return list(cuentas.values()), errores, avisos


def procesar_archivo(nombre: str, contenido: bytes) -> tuple[list[dict], list[str], list[str]]:
    """Devuelve (filas válidas, errores, avisos). Lanza ValueError si el archivo no sirve."""
    cruzado = _desde_coa_pcge(nombre, contenido)
    if cruzado is not None:
        return cruzado
    filas, errores = _procesar_simple(nombre, contenido)
    return filas, errores, []


def _procesar_simple(nombre: str, contenido: bytes) -> tuple[list[dict], list[str]]:
    filas = _filas_de_archivo(nombre, contenido)
    if not filas:
        raise ValueError("El archivo está vacío.")
    i_fa, i_ns, n_fa, n_ns = _columnas(filas[0])
    cuentas: dict[str, dict] = {}
    errores: list[str] = []
    for nro, f in enumerate(filas[1:], start=2):
        f = list(f) + [None] * 8
        falcon, ns = _texto(f[i_fa]), _texto(f[i_ns])
        if not falcon and not ns:
            continue
        if not falcon or not ns:
            errores.append(f"Fila {nro}: falta {'la cuenta Falcon' if not falcon else 'la cuenta NetSuite'}.")
            continue
        clave = normalizar(falcon)
        previa = cuentas.get(clave)
        if previa and previa["netsuite"] != ns:
            errores.append(f"Fila {nro}: la cuenta Falcon {falcon} ya estaba asignada a {previa['netsuite']} y ahora a {ns}.")
            continue
        cuentas[clave] = {
            "falcon": falcon, "netsuite": ns,
            "nombre_falcon": _texto(f[n_fa]) if n_fa is not None else "",
            "nombre_netsuite": _texto(f[n_ns]) if n_ns is not None else "",
        }
        if len(cuentas) > MAX_FILAS:
            raise ValueError(f"Demasiadas cuentas (máximo {MAX_FILAS}).")
    if not cuentas:
        raise ValueError("No se encontró ninguna cuenta con datos.")
    return list(cuentas.values()), errores


def guardar(filas: list[dict], archivo: str) -> None:
    RUTA.parent.mkdir(parents=True, exist_ok=True)
    RUTA.write_text(json.dumps({"archivo": archivo, "actualizado": time.strftime("%d/%m/%Y %H:%M"), "filas": filas},
                               ensure_ascii=False), encoding="utf-8")
    _cache.clear()


def borrar() -> None:
    RUTA.unlink(missing_ok=True)
    _cache.clear()


_cache: dict = {}


def cargar() -> dict:
    """{'archivo','actualizado','filas':[...]} (vacío si no hay tabla)."""
    try:
        mtime = RUTA.stat().st_mtime
    except FileNotFoundError:
        return {"archivo": "", "actualizado": "", "filas": []}
    if _cache.get("mtime") != mtime:
        try:
            _cache["datos"] = json.loads(RUTA.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            _cache["datos"] = {"archivo": "", "actualizado": "", "filas": []}
        _cache["mtime"] = mtime
        _cache["mapa"] = {normalizar(f["falcon"]): f["netsuite"] for f in _cache["datos"]["filas"]}
        _cache["coa"] = {normalizar(f["falcon"]): coa_de_nombre(f.get("nombre_netsuite", "")) for f in _cache["datos"]["filas"]}
    return _cache["datos"]


RUTA_NOMBRES_ES = Path(__file__).resolve().parent.parent / "data" / "nombres_pcge_es.json"


def nombres_es() -> dict[str, str]:
    """Nombre en español de las cuentas PCGE conocidas ({cuenta Falcon/PCGE: nombre}); el resto se muestra como viene en la tabla."""
    try:
        return json.loads(RUTA_NOMBRES_ES.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def coa_de_nombre(nombre: str) -> str:
    """Número de la cuenta en el COA de NetSuite: primera palabra del nombre si es numérica (p. ej. «9011.06 FCP-Related Costs» -> 9011.06)."""
    primero = (nombre or "").strip().split(" ", 1)[0]
    return primero if re.fullmatch(r"\d+(\.\d+)*", primero) else ""


def coa() -> dict[str, str]:
    """{cuenta Falcon/PCGE: número de cuenta del COA de NetSuite}."""
    cargar()
    return _cache.get("coa", {})


def mapa() -> dict[str, str]:
    cargar()
    return _cache.get("mapa", {})


def plantilla_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuentas homologadas"
    ws.append(["Cuenta Falcon", "Nombre Falcon", "ID interno NetSuite", "Nombre NetSuite"])
    for col, ancho in zip("ABCD", (18, 38, 18, 38)):
        ws.column_dimensions[col].width = ancho
    ws2 = wb.create_sheet("Instrucciones")
    for linea in (
        "Una fila por cuenta. Borre esta hoja o déjela: solo se lee la primera hoja.",
        "Cuenta Falcon: código contable que usa en Falcon Contable (por ejemplo 1041).",
        "ID interno NetSuite: el ID interno de la cuenta en NetSuite (lo que se escribirá en el CSV).",
        "Los nombres son opcionales y solo sirven para que usted revise la tabla.",
        "Una misma cuenta Falcon no puede tener dos cuentas NetSuite distintas.",
    ):
        ws2.append([linea])
    ws2.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
